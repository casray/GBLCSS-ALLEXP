import hashlib
import io
import json
import posixpath
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree as ET
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject
import pypdfium2 as pdfium


SOURCE = Path(r"C:\Users\Admin\Downloads\final.zip")
WORK = Path(__file__).resolve().parent
OUTPUT = WORK.parent.parent / "final_metadata_clean.zip"
STAMP = (1980, 1, 1, 0, 0, 0)
PDF_KEYS = {"/Metadata", "/PieceInfo", "/LastModified", "/Thumb"}
XML_META = {"rsids", "rsid", "rsidRoot", "docId", "attachedTemplate", "savePreviewPicture"}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def add_entry(archive, name, data):
    info = zipfile.ZipInfo(name, STAMP)
    info.create_system = 0
    info.external_attr = 16 if name.endswith("/") else 32
    info.compress_type = zipfile.ZIP_DEFLATED
    archive.writestr(info, data)


def scrub_pdf_dicts(objects, change):
    found = []
    seen = set()

    def walk(obj):
        if id(obj) in seen:
            return
        seen.add(id(obj))
        if isinstance(obj, DictionaryObject):
            for key in list(obj):
                if key in PDF_KEYS or key.startswith("/PTEX."):
                    found.append(str(key))
                    if change:
                        del obj[key]
                else:
                    walk(obj.raw_get(key))
        elif isinstance(obj, ArrayObject):
            for child in obj:
                walk(child)

    for obj in objects:
        walk(obj)
    return sorted(set(found))


def clean_pdf(data, name):
    before = PdfReader(io.BytesIO(data))
    writer = PdfWriter(clone_from=io.BytesIO(data))
    info_keys = sorted((before.metadata or {}).keys())
    # Clear the dictionary as well as its trailer reference, then prune orphans.
    if writer.metadata is not None:
        writer.metadata.clear()
    writer.metadata = None
    writer._ID = None
    removed = scrub_pdf_dicts(writer._objects, True)
    writer.compress_identical_objects(remove_duplicates=False, remove_unreferenced=True)
    stream = io.BytesIO()
    writer.write(stream)
    result = stream.getvalue()
    after = PdfReader(io.BytesIO(result))
    assert after.metadata is None and "/ID" not in after.trailer
    check = PdfWriter(clone_from=io.BytesIO(result))
    assert not scrub_pdf_dicts(check._objects, False)
    assert b"%3Cmxfile" not in result and b"/PTEX." not in result
    assert len(before.pages) == len(after.pages)
    assert [p.extract_text() for p in before.pages] == [p.extract_text() for p in after.pages]
    original = pdfium.PdfDocument(data)
    cleaned = pdfium.PdfDocument(result)
    for index in range(len(original)):
        pa, pb = original[index], cleaned[index]
        a, b = pa.render(scale=1.5), pb.render(scale=1.5)
        ia, ib = a.to_pil(), b.to_pil()
        assert ia.size == ib.size and ia.tobytes() == ib.tobytes(), (name, index)
        if name == "main.pdf" and index == 0:
            ib.save(WORK / "main_clean_preview.png")
        ia.close()
        ib.close()
        a.close()
        b.close()
        pa.close()
        pb.close()
    original.close()
    cleaned.close()
    return result, {"info_fields_removed": info_keys, "other_fields_removed": removed,
                    "pages": len(before.pages), "text_identical": True, "all_page_pixels_identical": True}


def strip_xml_meta(tree):
    changed = False
    for element in list(tree.iter()):
        if ET.QName(element).localname in XML_META:
            if element.getparent() is not None:
                element.getparent().remove(element)
                changed = True
            continue
        for attribute in list(element.attrib):
            local = ET.QName(attribute).localname
            if local.startswith("rsid") or local in {"paraId", "textId", "docId"}:
                del element.attrib[attribute]
                changed = True
    return changed


def canonical_content(data):
    tree = ET.fromstring(data)
    strip_xml_meta(tree)
    return ET.tostring(tree, method="c14n")


def clean_docx(data):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = {n: archive.read(n) for n in archive.namelist()}
    removed = {n for n in entries if n.startswith(("docProps/", "customXml/"))}
    document = ET.fromstring(entries["word/document.xml"])
    assert not document.xpath("//*[local-name()='dataBinding' or local-name()='ins' or local-name()='del']")
    output = {}
    for name, content in entries.items():
        if name in removed:
            continue
        if name.endswith((".xml", ".rels")):
            tree = ET.fromstring(content)
            changed = strip_xml_meta(tree)
            if name == "[Content_Types].xml":
                for child in list(tree):
                    if child.get("PartName", "").lstrip("/") in removed:
                        tree.remove(child)
                        changed = True
            if name.endswith(".rels"):
                parent = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
                for child in list(tree):
                    target = posixpath.normpath(posixpath.join(parent, child.get("Target", "")))
                    if target in removed:
                        tree.remove(child)
                        changed = True
            if changed:
                content = ET.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        output[name] = content
    for name in entries:
        if name.startswith("word/") and name.endswith(".xml"):
            assert canonical_content(entries[name]) == canonical_content(output[name]), name
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as cleaned:
        for name, content in output.items():
            add_entry(cleaned, name, content)
    result = stream.getvalue()
    a, b = Document(io.BytesIO(data)), Document(io.BytesIO(result))
    assert [p.text for p in a.paragraphs] == [p.text for p in b.paragraphs]
    with zipfile.ZipFile(io.BytesIO(result)) as check:
        assert check.testzip() is None
        assert not any(n.startswith(("docProps/", "customXml/")) for n in check.namelist())
        for entry in check.infolist():
            assert entry.date_time == STAMP and not entry.extra and not entry.comment
    return result, {"removed_parts": sorted(removed), "visible_content_and_formatting_xml_identical": True,
                    "paragraph_text_identical": True}


def main():
    source_hash = digest(SOURCE.read_bytes())
    report = {"source_sha256": source_hash, "pdf": {}, "docx": {}, "unchanged_content_files": []}
    temporary = OUTPUT.with_suffix(".pending.zip")
    with zipfile.ZipFile(SOURCE) as source, zipfile.ZipFile(temporary, "w") as target:
        assert source.testzip() is None
        for entry in source.infolist():
            name = entry.filename
            data = source.read(name)
            basename = posixpath.basename(name)
            if name.endswith(".pdf"):
                data, report["pdf"][basename] = clean_pdf(data, basename)
            elif name.endswith(".docx"):
                data, report["docx"][basename] = clean_docx(data)
            elif not entry.is_dir():
                report["unchanged_content_files"].append(basename)
            add_entry(target, name, data)
            print("Verified", basename or name, flush=True)
    with zipfile.ZipFile(SOURCE) as source, zipfile.ZipFile(temporary) as target:
        assert source.namelist() == target.namelist()
        assert target.testzip() is None and not target.comment
        for entry in target.infolist():
            assert entry.date_time == STAMP and not entry.extra and not entry.comment
            if not entry.filename.endswith((".pdf", ".docx")):
                assert source.read(entry.filename) == target.read(entry.filename)
    assert source_hash == digest(SOURCE.read_bytes())
    temporary.replace(OUTPUT)
    report["output_sha256"] = digest(OUTPUT.read_bytes())
    report["output_path"] = str(OUTPUT)
    report["source_unchanged"] = True
    (WORK / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print("OUTPUT", OUTPUT, flush=True)
    print("PDF_PAGES_VERIFIED", sum(x["pages"] for x in report["pdf"].values()), flush=True)


if __name__ == "__main__":
    main()
