from copy import deepcopy
from pathlib import Path
import hashlib
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path('C:/Users/Admin/Downloads/PR_Cover_Letter.docx')
OUTPUT = ROOT / 'PR_Cover_Letter_polished_20260930.docx'
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
assert len(doc.paragraphs) == 15

REPLACEMENTS = {
    5: (
        'Granular Ball Local-Consensus Soft Supervision (GB-LCSS) addresses a practical '
        'problem: learning from noisy node labels when graph '
        'connections are available during training but cannot be used at deployment. Its '
        'central contribution is to use granular balls to organize supervision. The balls '
        'are formed geometrically in a frozen, label-free graph representation; only then '
        'are other training nodes\' labels within each ball used to refine a node\'s target. '
        'This separation keeps noisy labels out of region construction and transfers graph '
        'information into supervision for an attribute-only multilayer perceptron (MLP).'
    ),
    7: (
        'Across five datasets and five noise settings, GB-LCSS achieves the highest '
        'five-seed mean accuracy in 24 of 25 settings against seven noisy-label baselines '
        'with the same deployed two-layer MLP. Its average gain over JoCoR, the strongest '
        'baseline on average, is 3.46 percentage points. The refined targets also have '
        'lower mean Brier scores than the observed one-hot labels in all 25 settings. Together '
        'with the local-vote analysis and grouping ablations, these results support the '
        'value of improving supervision while keeping the deployed classifier unchanged.'
    ),
    8: (
        'We believe this combination of robust classification, graph representation '
        'learning, and granular-ball modeling will interest readers of Pattern Recognition, '
        'particularly those developing learning methods for unreliable annotations and '
        'constrained deployment.'
    ),
}


def replace_text(paragraph, text):
    fmt = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if fmt is not None:
        run._r.insert(0, fmt)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)


for index, text in REPLACEMENTS.items():
    replace_text(doc.paragraphs[index], text)

paragraph = doc.paragraphs[6]
paragraph.clear()
comparison_parts = [
    ('Our comparisons cover noisy-label learning, graph-assisted MLPs, and graph-to-MLP '
     'distillation. Five representative published methods are ', None),
    ('LogitClip (ICML 2023)', 'https://proceedings.mlr.press/v202/wei23e.html'),
    (', ', None),
    ('PosteL (AAAI 2026)', 'https://ojs.aaai.org/index.php/AAAI/article/view/39321'),
    (', ', None),
    ('FF-G2M (AAAI 2023)', 'https://ojs.aaai.org/index.php/AAAI/article/view/26232'),
    (', ', None),
    ('KRD (ICML 2023)', 'https://proceedings.mlr.press/v202/wu23m.html'),
    (', and ', None),
    ('VQGraph (ICLR 2024)', 'https://openreview.net/forum?id=h6Tz85BqRI'),
    ('. These methods were developed for different tasks and backbones. We evaluate '
     'them under shared data splits and noise realizations, adapting their mechanisms '
     'where necessary. The Supplementary documents these adaptations. The comparisons '
     'assess robustness under our common deployment protocol, rather than establish '
     'an exhaustive state-of-the-art ranking.', None),
]
for text, url in comparison_parts:
    if url is None:
        run = paragraph.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
    else:
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Times New Roman')
        fonts.set(qn('w:hAnsi'), 'Times New Roman')
        props.append(fonts)
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '22')
        props.append(size)
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        props.append(color)
        run.append(props)
        content = OxmlElement('w:t')
        content.text = text
        run.append(content)
        link.append(run)
        paragraph._p.append(link)

doc.core_properties.title = 'Cover letter for Pattern Recognition'
doc.save(OUTPUT)

check = Document(OUTPUT)
texts = [''.join(p._p.xpath('.//w:t/text()')) for p in check.paragraphs]
original = Document(SOURCE)
for index in [0, 1, 2, 3, 4, 9, 10, 11, 12, 13, 14]:
    assert texts[index] == original.paragraphs[index].text, index
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
all_text = '\n'.join(texts)
assert all(name in all_text for name in ['LogitClip', 'PosteL', 'FF-G2M', 'KRD', 'VQGraph'])
assert '3.46' in all_text and '24 of 25' in all_text
assert 'No existing method' not in all_text
print(json.dumps({
    'output': str(OUTPUT),
    'words_including_heading_signature': len(all_text.split()),
    'characters_including_newlines': len(all_text),
    'source_unchanged': True,
    'body_font_points': 11,
    'changed_paragraphs': [5, 6, 7, 8],
}, indent=2))
