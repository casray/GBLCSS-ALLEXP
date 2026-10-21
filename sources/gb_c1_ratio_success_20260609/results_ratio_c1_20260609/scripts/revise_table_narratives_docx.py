from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


BODY_REPLACEMENTS = {
    "表 2 给出采用相同部署 MLP 架构的抗噪方法在 uniform 0.4 下的比较结果。":
        "表 2 首先检验局部共识监督相对于同部署架构抗噪方法的效果。",
    "表 3 比较结构辅助 X-only 方法在 uniform 0.4 下的表现。":
        "表 3 将比较范围扩展到训练阶段利用图信息的 X-only 方法。",
    "表 4 给出不同图知识迁移路线在 uniform 0.4 和 0.6 下的比较结果。":
        "表 4 进一步比较基础 GB-LCSS 与不同图知识迁移路线。",
}

TABLE4_OLD = "表 4 基础 GB-LCSS 与五条图知识迁移路线的比较。"
TABLE4_REST = " 基础 GB-LCSS 与五条图知识迁移路线在 uniform 0.4 和 0.6 下的测试精度。"


def replace_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    doc = Document(args.input)
    replaced = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in BODY_REPLACEMENTS:
            replace_paragraph_text(paragraph, BODY_REPLACEMENTS[text])
            replaced.append(text)
        elif text == TABLE4_OLD:
            paragraph.clear()
            paragraph.add_run("表 4").bold = True
            paragraph.add_run(TABLE4_REST)
            replaced.append(text)

    expected = len(BODY_REPLACEMENTS) + 1
    if len(replaced) != expected:
        raise RuntimeError(f"Expected {expected} replacements, found {len(replaced)}")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"saved={args.output}")
    print(f"replacements={len(replaced)}")


if __name__ == "__main__":
    main()
