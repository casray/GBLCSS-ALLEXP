from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "23405F"
MID_BLUE = "DCE6F1"
PALE_BLUE = "EEF3F8"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "AEB7C1"
TEXT_GRAY = "4B5563"
WHITE = "FFFFFF"


def set_east_asia_font(run_or_style, east_asia: str, latin: str | None = None) -> None:
    font = run_or_style.font
    latin = latin or east_asia
    font.name = latin
    rpr = run_or_style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size=5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_fixed_layout(table, total_twips: int) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_style(doc: Document, name: str, east_asia: str, latin: str, size: float,
              color: str | None = None, bold: bool | None = None) -> None:
    style = doc.styles[name]
    set_east_asia_font(style, east_asia, latin)
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def configure_styles(doc: Document) -> None:
    set_style(doc, "Normal", "宋体", "Times New Roman", 10.5, "222222")
    normal = doc.styles["Normal"].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.line_spacing = 1.25
    normal.space_after = Pt(5)
    normal.first_line_indent = Cm(0.74)

    for name in ("Body Text", "First Paragraph"):
        set_style(doc, name, "宋体", "Times New Roman", 10.5, "222222")
        pf = doc.styles[name].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.25
        pf.space_after = Pt(5)
        pf.first_line_indent = Cm(0.74)

    if "Compact" in [s.name for s in doc.styles]:
        set_style(doc, "Compact", "宋体", "Times New Roman", 10.5, "222222")
        pf = doc.styles["Compact"].paragraph_format
        pf.line_spacing = 1.18
        pf.space_after = Pt(2)

    heading_specs = {
        "Heading 1": (15.5, 10, 5),
        "Heading 2": (14.0, 9, 4),
        "Heading 3": (12.0, 7, 3),
        "Heading 4": (11.0, 6, 2),
    }
    for name, (size, before, after) in heading_specs.items():
        set_style(doc, name, "黑体", "Arial", size, NAVY, True)
        pf = doc.styles[name].paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.line_spacing = 1.05
        pf.first_line_indent = Cm(0)

    if "Caption" in [s.name for s in doc.styles]:
        set_style(doc, "Caption", "宋体", "Times New Roman", 9.0, TEXT_GRAY)
        pf = doc.styles["Caption"].paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.1
        pf.space_before = Pt(3)
        pf.space_after = Pt(7)
        pf.first_line_indent = Cm(0)

    if "Source Code" in [s.name for s in doc.styles]:
        set_style(doc, "Source Code", "宋体", "Cambria Math", 9.5, "222222")


def configure_page(doc: Document) -> int:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.05)
        section.right_margin = Cm(2.05)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        header = section.header
        p = header.paragraphs[0]
        p.text = "粒球局部共识软监督 · 导师审阅稿"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            set_east_asia_font(run, "宋体", "Times New Roman")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("7A8592")

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run()
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        run._r.extend([fld_char1, instr, fld_char2])
        set_east_asia_font(run, "宋体", "Times New Roman")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("7A8592")

    usable_cm = 21.0 - 2.05 - 2.05
    return int(usable_cm / 2.54 * 1440)


def style_title_block(doc: Document) -> None:
    if not doc.paragraphs:
        return
    title = doc.paragraphs[0]
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(5)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        set_east_asia_font(run, "黑体", "Arial")
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)

    if len(doc.paragraphs) > 1:
        subtitle = doc.paragraphs[1]
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(11)
        subtitle.paragraph_format.first_line_indent = Cm(0)
        for run in subtitle.runs:
            set_east_asia_font(run, "宋体", "Times New Roman")
            run.font.size = Pt(10.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def body_children(parent):
    return list(parent.element.body.iterchildren())


def mark_captions(doc: Document) -> None:
    body = doc.element.body
    children = list(body.iterchildren())
    for idx, child in enumerate(children):
        if child.tag == qn("w:tbl") and idx > 0:
            prev = children[idx - 1]
            if prev.tag == qn("w:p"):
                p_text = "".join(prev.itertext()).strip()
                if re.match(r"^(?:表|Table)\s+[A-Z]?\d+(?:\.\d+)?\.?\b", p_text):
                    from docx.text.paragraph import Paragraph
                    p = Paragraph(prev, doc)
                    p.style = doc.styles["Caption"]
        if child.tag == qn("w:p") and child.xpath(".//w:drawing") and idx + 1 < len(children):
            nxt = children[idx + 1]
            if nxt.tag == qn("w:p"):
                from docx.text.paragraph import Paragraph
                p = Paragraph(nxt, doc)
                if re.match(r"^(?:图|Fig\.)\s+\d+\b", p.text.strip()):
                    p.style = doc.styles["Caption"]


def add_algorithm_table_before(doc: Document, code_paragraph, lines: list[str]):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table"
    set_table_fixed_layout(table, 9340)
    set_table_borders(table, color="7E8B99", size=6)

    parsed = []
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^\s*(\d+):\s?(.*)$", line)
        if m:
            parsed.append((m.group(1), m.group(2), len(line) - len(line.lstrip())))
        elif line.startswith("输入:") or line.startswith("输出:"):
            parsed.append(("", line, 0))
        else:
            parsed.append(("", line.strip(), 0))

    for row_idx, (number, text, indent) in enumerate(parsed):
        row = table.add_row()
        prevent_row_split(row)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_width(row.cells[0], 640)
        set_cell_width(row.cells[1], 8700)
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45, start=75, end=75)
        p_num = row.cells[0].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_num.paragraph_format.space_after = Pt(0)
        p_text = row.cells[1].paragraphs[0]
        p_text.paragraph_format.space_after = Pt(0)
        p_text.paragraph_format.line_spacing = 1.05
        if indent:
            p_text.paragraph_format.left_indent = Cm(min(indent, 12) * 0.08)
        r1 = p_num.add_run(number)
        r2 = p_text.add_run(text)
        for run in (r1, r2):
            set_east_asia_font(run, "宋体", "Cambria Math")
            run.font.size = Pt(9.25)
        if text.startswith("输入:") or text.startswith("输出:"):
            shade_cell(row.cells[0], PALE_BLUE)
            shade_cell(row.cells[1], PALE_BLUE)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor.from_string(NAVY)
        elif row_idx % 2 == 1:
            shade_cell(row.cells[0], "FAFBFC")
            shade_cell(row.cells[1], "FAFBFC")

    code_paragraph._p.addprevious(table._tbl)
    code_paragraph._element.getparent().remove(code_paragraph._element)
    return table


def convert_algorithms(doc: Document) -> list:
    algorithm_tables = []
    for p in list(doc.paragraphs):
        if p.text.strip().startswith(("算法 ", "Algorithm ")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                set_east_asia_font(run, "黑体", "Arial")
                run.font.size = Pt(10.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(NAVY)
            nxt = p._p.getnext()
            if nxt is not None and nxt.tag == qn("w:p"):
                from docx.text.paragraph import Paragraph
                code = Paragraph(nxt, doc)
                if code.style.name == "Source Code":
                    algorithm_tables.append(add_algorithm_table_before(doc, code, code.text.splitlines()))
    return algorithm_tables


def text_width_score(text: str) -> float:
    score = 0.0
    for ch in text:
        if "\u4e00" <= ch <= "\u9fff":
            score += 1.7
        elif ch in "MW@%":
            score += 1.3
        else:
            score += 1.0
    return score


def table_column_widths(table, total_twips: int) -> list[int]:
    ncols = len(table.columns)
    weights = []
    for col_idx in range(ncols):
        samples = [row.cells[col_idx].text.replace("\n", " ") for row in table.rows[: min(20, len(table.rows))]]
        max_score = max([text_width_score(x) for x in samples] + [5.0])
        if col_idx == 0:
            max_score = max(max_score, 12.0)
        weights.append(min(max_score, 30.0))
    if ncols >= 6:
        weights = [max(w, 8.0) for w in weights]
    total = sum(weights)
    widths = [max(720, int(total_twips * w / total)) for w in weights]
    correction = total_twips - sum(widths)
    widths[-1] += correction
    return widths


def style_tables(doc: Document, usable_twips: int, algorithm_tables: list) -> None:
    algorithm_ids = {id(t._tbl) for t in algorithm_tables}
    for table in doc.tables:
        if id(table._tbl) in algorithm_ids:
            continue
        ncols = len(table.columns)
        font_size = 8.7 if ncols <= 6 else (8.0 if ncols == 7 else 7.3)
        if len(table.rows) > 40:
            font_size = min(font_size, 7.8)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table"
        set_table_fixed_layout(table, usable_twips)
        set_table_borders(table, color="B8C1CB", size=4)
        widths = table_column_widths(table, usable_twips)
        for row_idx, row in enumerate(table.rows):
            prevent_row_split(row)
            if row_idx == 0:
                repeat_table_header(row)
            for col_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_width(cell, widths[col_idx])
                set_cell_margins(cell, top=55, bottom=55, start=60, end=60)
                if row_idx == 0:
                    shade_cell(cell, MID_BLUE)
                elif row_idx % 2 == 0:
                    shade_cell(cell, "F8FAFC")
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.first_line_indent = Cm(0)
                    if col_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        set_east_asia_font(run, "宋体", "Times New Roman")
                        run.font.size = Pt(font_size)
                        if row_idx == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string(NAVY)


def style_paragraphs(doc: Document) -> None:
    in_references = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "参考文献":
            in_references = True
        elif p.style.name.startswith("Heading") and text != "参考文献":
            in_references = False

        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            continue

        if p.style.name == "Caption":
            for run in p.runs:
                set_east_asia_font(run, "宋体", "Times New Roman")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
            continue

        if p.style.name.startswith("Heading"):
            continue

        if text.startswith(("关键词", "Keywords")):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                set_east_asia_font(run, "宋体", "Times New Roman")
                run.font.size = Pt(10.5)
            continue

        num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
        if num_pr is not None:
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_after = Pt(3)

        if in_references and text:
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_east_asia_font(run, "宋体", "Times New Roman")
                run.font.size = Pt(9)

        if text.startswith("注："):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.right_indent = Cm(0.35)
            for run in p.runs:
                set_east_asia_font(run, "宋体", "Times New Roman")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def resize_images(doc: Document, max_width_inches: float = 6.35) -> None:
    max_width = Inches(max_width_inches)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)


def extract_display_equation_tags(source_path: str | None) -> list[str | None]:
    if not source_path:
        return []
    lines = Path(source_path).read_text(encoding="utf-8").splitlines()
    tags: list[str | None] = []
    i = 0
    while i < len(lines):
        if lines[i].strip() == "$$":
            j = i + 1
            body: list[str] = []
            while j < len(lines) and lines[j].strip() != "$$":
                body.append(lines[j])
                j += 1
            if j >= len(lines):
                raise RuntimeError(f"Unclosed display-math block at line {i + 1}")
            match = re.search(r"\\tag\{([^}]+)\}", "\n".join(body))
            tags.append(match.group(1) if match else None)
            i = j + 1
        else:
            i += 1
    return tags


def set_paragraph_tab(p_pr, kind: str, pos: int) -> None:
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), kind)
    tab.set(qn("w:pos"), str(pos))
    tabs.append(tab)


def tab_run() -> OxmlElement:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    return run


def equation_number_run(tag: str) -> OxmlElement:
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{attr}"), "Times New Roman")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = f"({tag})"
    run.append(text)
    return run


def number_display_equations(doc: Document, tags: list[str | None], usable_twips: int) -> None:
    equations = [p for p in doc.paragraphs if p._p.xpath(".//m:oMathPara")]
    if tags and len(tags) != len(equations):
        raise RuntimeError(
            f"Display-equation mismatch: source has {len(tags)}, DOCX has {len(equations)}"
        )
    for p, tag in zip(equations, tags):
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.right_indent = Cm(0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        math_paras = p._p.xpath("./m:oMathPara")
        if not math_paras:
            continue
        math_para = math_paras[0]
        math_node = math_para.find(qn("m:oMath"))
        if math_node is None:
            continue
        math_para.remove(math_node)
        p._p.replace(math_para, math_node)

        if tag:
            p_pr = p._p.get_or_add_pPr()
            set_paragraph_tab(p_pr, "center", usable_twips // 2)
            set_paragraph_tab(p_pr, "right", usable_twips)
            math_index = list(p._p).index(math_node)
            p._p.insert(math_index, tab_run())
            p._p.append(tab_run())
            p._p.append(equation_number_run(tag))
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_breaks(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "补充材料":
            p.paragraph_format.page_break_before = True


def clean_properties(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = "导师审阅稿"
    props.keywords = "GB-LCSS; noisy-label learning; X-only inference"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--source-markdown")
    args = parser.parse_args()

    doc = Document(args.input)
    configure_styles(doc)
    usable_twips = configure_page(doc)
    style_title_block(doc)
    algorithm_tables = convert_algorithms(doc)
    equation_tags = extract_display_equation_tags(args.source_markdown)
    number_display_equations(doc, equation_tags, usable_twips)
    mark_captions(doc)
    style_tables(doc, usable_twips, algorithm_tables)
    style_paragraphs(doc)
    resize_images(doc)
    add_page_breaks(doc)
    clean_properties(doc)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"saved={output}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")
    print(f"omath={doc._element.xml.count('<m:oMath')}")


if __name__ == "__main__":
    main()
