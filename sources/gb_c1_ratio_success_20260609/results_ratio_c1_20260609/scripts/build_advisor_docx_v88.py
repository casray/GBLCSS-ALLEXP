from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


VML_RECT = "{urn:schemas-microsoft-com:vml}rect"


def remove_horizontal_rules(doc: Document) -> int:
    removed = 0
    for paragraph in list(doc.paragraphs):
        if any(element.tag == VML_RECT for element in paragraph._p.iter()):
            parent = paragraph._p.getparent()
            if parent is not None:
                parent.remove(paragraph._p)
                removed += 1
    return removed


def shade_cell_white(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "FFFFFF")


def set_table_borders(table, color: str = "666666", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def style_tables_white(doc: Document) -> None:
    for table in doc.tables:
        set_table_borders(table)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                shade_cell_white(cell)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        color = run._r.get_or_add_rPr().find(qn("w:color"))
                        if color is None:
                            color = OxmlElement("w:color")
                            run._r.get_or_add_rPr().append(color)
                        color.set(qn("w:val"), "000000")
                        if row_index == 0:
                            run.bold = True


def set_table_column_widths(table, widths_cm: list[float]) -> None:
    if len(table.columns) != len(widths_cm):
        raise RuntimeError("Column-width specification does not match the table")

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    widths_twips = [round(width / 2.54 * 1440) for width in widths_cm]
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")

    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_column, width_twips in zip(grid_columns, widths_twips):
        grid_column.set(qn("w:w"), str(width_twips))

    for column_index, width_twips in enumerate(widths_twips):
        for row in table.rows:
            tc_pr = row.cells[column_index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")


def adjust_structure_assisted_table(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows or len(table.columns) != 7:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:2] == ["Method", "Graph information used in training"]:
            set_table_column_widths(table, [1.65, 2.95, 2.38, 2.38, 2.38, 2.38, 2.38])
            return


def insert_algorithm_image(
    doc: Document,
    caption_prefix: str,
    image_path: Path,
    *,
    page_break_before: bool,
) -> None:
    caption = next(
        (paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith(caption_prefix)),
        None,
    )
    if caption is None:
        raise RuntimeError(f"Algorithm caption not found: {caption_prefix}")

    following = caption._p.getnext()
    if following is None or following.tag != qn("w:tbl"):
        raise RuntimeError(f"Algorithm table not found after: {caption_prefix}")

    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Cm(0)
    image_paragraph.paragraph_format.left_indent = Cm(0)
    image_paragraph.paragraph_format.right_indent = Cm(0)
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(6)
    image_paragraph.paragraph_format.page_break_before = page_break_before
    image_paragraph.add_run().add_picture(str(image_path), width=Cm(16.5))

    caption._p.addprevious(image_paragraph._p)
    following.getparent().remove(following)
    caption._p.getparent().remove(caption._p)


def clean_properties(doc: Document) -> None:
    props = doc.core_properties
    props.author = "Anonymous"
    props.last_modified_by = "Anonymous"
    props.comments = ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--algorithm1", required=True, type=Path)
    parser.add_argument("--algorithm2", required=True, type=Path)
    parser.add_argument("--caption1", default="算法 1")
    parser.add_argument("--caption2", default="算法 2")
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    doc = Document(args.input)
    removed_rules = remove_horizontal_rules(doc)
    insert_algorithm_image(doc, args.caption1, args.algorithm1, page_break_before=True)
    insert_algorithm_image(doc, args.caption2, args.algorithm2, page_break_before=False)
    style_tables_white(doc)
    adjust_structure_assisted_table(doc)
    clean_properties(doc)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"saved={args.output}")
    print(f"horizontal_rules_removed={removed_rules}")
    print(f"tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
