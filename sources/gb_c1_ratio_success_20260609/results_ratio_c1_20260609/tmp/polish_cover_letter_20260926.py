from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[1]
folder = root / 'GB-LCSS_EM_submission_20260926'
source = folder / 'cover_letter.docx'
document = Document(source)

replacements = {
    3: (
        'Please consider our manuscript for publication as a Research Article in '
        'Information Sciences. We address learning from noisy node labels when '
        'graph structure is available during training but prediction must rely '
        'on node attributes alone.'
    ),
    4: (
        'The central contribution of Granular Ball Local-Consensus Soft Supervision '
        '(GB-LCSS) is to use the training graph to refine noisy supervision for an '
        'attribute-only classifier. The method separates label-free region '
        'construction from noisy-label aggregation. Granular balls organize '
        'shared local label statistics, which are then used to construct '
        'node-specific soft targets.'
    ),
    5: (
        "A leave-one-out vote excludes each node's annotation from its own "
        'consensus estimate. This design also supports a conditional error '
        'analysis: under independent symmetric label noise, we derive a '
        'sufficient condition for the soft target to have lower expected squared '
        'error than the observed one-hot label. Controlled experiments assess '
        'the roles of local grouping and target construction.'
    ),
    6: (
        'The manuscript connects granular ball computing with robust learning '
        'on graph-structured data. We believe this contribution will be of '
        'interest to Information Sciences readers working on granular computing, '
        'intelligent classification, and learning with noisy labels.'
    ),
}

assert len(document.paragraphs) == 10
for index, text in replacements.items():
    paragraph = document.paragraphs[index]
    assert len(paragraph.runs) == 1, index
    paragraph.runs[0].text = text

output = folder / 'cover_letter_polished.docx'
document.save(output)
check = Document(output)
original = Document(source)
for index in range(len(original.paragraphs)):
    expected = replacements.get(index, original.paragraphs[index].text)
    assert check.paragraphs[index].text == expected
assert check.sections[0]._sectPr.xml == original.sections[0]._sectPr.xml
print(output)
print('Body word count:', sum(len(p.split()) for p in replacements.values()))
