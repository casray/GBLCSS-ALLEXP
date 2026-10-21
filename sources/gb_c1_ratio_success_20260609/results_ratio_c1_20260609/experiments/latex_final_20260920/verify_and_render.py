from pathlib import Path
import json
import re
import hashlib
from collections import Counter

from docx import Document
import pypdfium2 as pdfium
import pdfplumber
from PIL import Image, ImageOps, ImageDraw

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
BASE = ROOT / "GB-LCSS_Elsevier_LaTeX_discussion_layout_final_20260918"
OUT = ROOT / "GB-LCSS_Elsevier_LaTeX_five_dataset_final_20260920"
MAIN_WORD = Path(r"C:\Users\Admin\Downloads\GB-LCSS_five_dataset_experiments_revised_20260920_red.docx")
SUPP_WORD = Path(r"C:\Users\Admin\Downloads\GB-LCSS_supplementary_five_dataset_experiments_revised_20260920_red.docx")


def read(path):
    return path.read_text(encoding="utf-8-sig")


def same_bytes(a, b):
    return hashlib.sha256(a.read_bytes()).digest() == hashlib.sha256(b.read_bytes()).digest()


main = read(OUT / "main.tex")
old = read(BASE / "main.tex")
assert re.findall(r"\\cite\{([^}]+)\}", main) == re.findall(r"\\cite\{([^}]+)\}", old)
assert same_bytes(OUT / "references.bib", BASE / "references.bib")
assert same_bytes(OUT / "gblcss-numeric.bst", BASE / "gblcss-numeric.bst")
for path in (BASE / "figures").glob("*.pdf"):
    assert same_bytes(path, OUT / "figures" / path.name)
for path in (BASE / "algorithms").glob("*.tex"):
    assert same_bytes(path, OUT / "algorithms" / path.name)
for path in (BASE / "tables").glob("*.tex"):
    if path.name not in {"table_6.tex", "table_F_2.tex"}:
        normalize_layout = lambda t: re.sub(r"\\Needspace\{\d+pt\}\s*", "", t).replace("[!htb]", "[htbp]").replace(r"\\*", r"\\").replace(r"\arraystretch}{1.12}", r"\arraystretch}{1.18}")
        assert normalize_layout(read(path)) == normalize_layout(read(OUT / "tables" / path.name))
equations = dict(re.findall(r"\\begin\{equation\}\\label\{eq:(\d+)\}(.*?)\\end\{equation\}", main, re.S))
old_equations = dict(re.findall(r"\\begin\{equation\}\\label\{eq:(\d+)\}(.*?)\\end\{equation\}", old, re.S))
assert list(map(int, equations)) == list(range(1, 24))
for num, formula in old_equations.items():
    new_num = str(int(num) if int(num) < 14 else int(num) + 3)
    assert formula == equations[new_num], (num, new_num)

# Compare table body numbers against the approved Word, independently of CSV generation.
def numbers(text):
    return Counter(re.findall(r"(?<![A-Za-z])\d+(?:\.\d+)?", text.replace("A2X", "A²X")))


def body_numbers(path, skip_first=False):
    text = read(path).split(r"\endhead")[-1] if "table_I_1" in path.name else read(path)
    lines = []
    for line in text.splitlines():
        if " & " in line and any("\\pm" in line or name in line for name in ["Cora &", "CiteSeer &", "CoraFull &", "Coauthor-CS &", "Amazon-Photo &"]):
            lines.append(line.split(" & ", 1)[1] if skip_first else line)
    return numbers(" ".join(lines))


main_tables = Document(MAIN_WORD).tables
rep_table = next(t for t in main_tables if "GRACE" in t.rows[0].cells[1].text and len(t.columns) == 7)
assert numbers(" ".join(c.text for r in rep_table.rows[1:] for c in r.cells)) == body_numbers(OUT / "tables/table_6.tex")
supp_tables = Document(SUPP_WORD).tables
for table, name in zip(supp_tables[-2:], ["table_I_1.tex", "table_I_2.tex"]):
    expected = numbers(" ".join(c.text for r in table.rows[1:] for c in r.cells[1:]))
    actual = body_numbers(OUT / "tables" / name, skip_first=True)
    assert expected == actual, (name, expected - actual, actual - expected)

report = {"source_checks": "passed", "unchanged": ["main comparison tables", "existing equations", "algorithms", "five vector figures", "74-reference bibliography and citation sequence"], "word_table_checks": "Table 7, I.1, I.2 match"}

for stem in ["main", "supplementary"]:
    log = read(OUT / f"{stem}.log")
    errors = re.findall(r"^.*(?:Undefined control sequence|LaTeX Error|undefined references|multiply defined|Missing character).*$", log, re.M)
    assert not errors, errors
    with pdfplumber.open(OUT / f"{stem}.pdf") as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        assert "??" not in text
        (WORK / f"{stem}_pdf_text.txt").write_text(text, encoding="utf-8")
        page_stats = []
        for i, page in enumerate(pdf.pages):
            blocks = list(page.chars) + list(page.images)
            # Ignore page numbers and running footer when locating substantive content.
            blocks = [b for b in blocks if b.get("top", 0) < page.height - 72 and b.get("bottom", 0) > 45]
            occupied = [False] * int(page.height)
            for b in blocks:
                for y in range(max(45, int(b["top"])), min(int(page.height - 55), int(b["bottom"]) + 1)):
                    occupied[y] = True
            used = [y for y, value in enumerate(occupied) if value]
            max_gap = 0
            if used:
                max_gap = max((b - a - 1 for a, b in zip(used, used[1:])), default=0)
            page_stats.append({"page": i+1, "bottom": max(used) if used else None, "largest_internal_gap": max_gap})
        report[stem] = {"pages": len(pdf.pages), "layout": page_stats}
    pages_dir = WORK / f"{stem}_pages"
    pages_dir.mkdir(exist_ok=True)
    pdf = pdfium.PdfDocument(str(OUT / f"{stem}.pdf"))
    thumbs = []
    for i in range(len(pdf)):
        page = pdf[i]
        bitmap = page.render(scale=1.6)
        im = bitmap.to_pil().convert("RGB")
        im.save(pages_dir / f"page_{i+1:02d}.png")
        thumb = ImageOps.contain(im, (304, 398))
        panel = Image.new("RGB", (326, 430), "#eeeeee")
        panel.paste(thumb, ((326-thumb.width)//2, 20))
        ImageDraw.Draw(panel).text((12, 412), f"{stem} {i+1}", fill="black")
        thumbs.append(panel)
        page.close()
    pdf.close()
    for group in range(0, len(thumbs), 12):
        part = thumbs[group:group+12]
        contact = Image.new("RGB", (326*4, 430*((len(part)+3)//4)), "#dddddd")
        for j, thumb in enumerate(part):
            contact.paste(thumb, ((j%4)*326, (j//4)*430))
        contact.save(WORK / f"{stem}_contact_{group//12+1}.png")

(WORK / "checks.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
print(json.dumps(report, indent=2))
