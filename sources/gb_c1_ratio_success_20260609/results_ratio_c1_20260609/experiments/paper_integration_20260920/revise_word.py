"""Integrate the approved theory and experiment edits in two new Word files."""
from collections import Counter
from copy import deepcopy
import hashlib
import importlib.util
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
DOWNLOADS = Path('C:/Users/Admin/Downloads')
spec = importlib.util.spec_from_file_location('polish_helpers', HERE.parent / 'paper_polish_20260920/polish_word.py')
helpers = importlib.util.module_from_spec(spec)
spec.loader.exec_module(helpers)
text = helpers.all_text
find = helpers.find
rewrite = helpers.rewrite
color = helpers.color_element
EQ_MAP = {14: 17, 15: 18, 16: 19, 17: 20, 18: 21, 19: 22, 20: 23, 21: 14, 22: 15, 23: 16}
SECTION_MAP = {'3.4': '3.5', '3.5': '3.6', '3.6': '3.4'}
CHANGES = []


def replace_in_paragraph(p, old, new):
    """Replace text across Word runs without flattening the surrounding mathematics."""
    nodes = p._p.xpath('.//w:t')
    joined = ''.join(t.text or '' for t in nodes)
    assert joined.count(old) == 1, (old, joined)
    start = joined.index(old)
    end = start + len(old)
    pos = 0
    inserted = False
    for node in nodes:
        original = node.text or ''
        next_pos = pos + len(original)
        if pos < end and next_pos > start:
            a = max(0, start - pos)
            b = min(len(original), end - pos)
            node.text = original[:a] + (new if not inserted else '') + original[b:]
            node.set(qn('xml:space'), 'preserve')
            inserted = True
            color(node.getparent())
        pos = next_pos
    CHANGES.append({'old': old, 'new': new})


def rename_clean_proportions(paragraphs):
    count = 0
    for p in paragraphs:
        for sub in list(p._p.xpath('.//m:sSub')):
            base = sub.find(qn('m:e'))
            if ''.join(t.text or '' for t in base.iter(qn('m:t'))) != '\u03c0':
                continue
            replacement = OxmlElement('m:sSubSup')
            new_base = deepcopy(base)
            tokens = list(new_base.iter(qn('m:t')))
            assert len(tokens) == 1
            tokens[0].text = 'p'
            replacement.append(new_base)
            replacement.append(deepcopy(sub.find(qn('m:sub'))))
            sup = OxmlElement('m:sup')
            sup.append(helpers.helper.math_run('clean', upright=True))
            replacement.append(sup)
            sub.getparent().replace(sub, replacement)
            color(replacement)
            count += 1
    return count


MAIN_REWRITES = {
    '(2) We develop consensus-weighted':
        '(2) We develop consensus-weighted soft supervision within these regions. Each training node combines its observed label with the labels of the other training nodes in its ball, with their agreement determining the contribution of local consensus. We analyze the reliability of local voting under independent symmetric label noise and derive an exact decomposition of soft-target error.',
    'Equation (23) holds for':
        'Equation (23) expresses the target error as a weighted average of the observed-label and local-consensus errors, minus a nonnegative term determined by their disagreement. The identity holds for the data-dependent weight and tie-handling rule of GB-LCSS without assuming a noise distribution. Soft targets can reduce error on corrupted labels while perturbing initially correct labels; their aggregate benefit depends on the balance between these effects. Appendix H.2 gives the derivation and a sufficient condition for lower conditional expected target error under symmetric noise. Section 4.3.1 and Appendix I.2 examine this balance empirically.',
    'The representation determines which':
        'The representation determines which nodes are grouped together and hence which labels contribute to each target. Table 7 compares GRACE with raw node attributes $X$ and two-step propagated attributes $A^2X$ on Cora, CiteSeer, and CoraFull under uniform 0.4 noise. All three conditions, including the default GB-LCSS configuration, were rerun with five paired seeds using the setup in Appendix I.1. Propagation uses the original adjacency matrix without added self-loops or adjacency normalization, followed by L2 normalization of the node vectors as in Eq. (3). Ball counts are reported as ranges across seeds.',
    'With the same splitting parameters':
        'With the same splitting parameters, raw $X$ and $A^2X$ both produced a single ball in all 15 runs. Each node therefore drew its consensus from all other training nodes rather than a local subset. The identical partitions produced identical soft targets and classification accuracy for each paired seed. Under the same splitting rule, GRACE produced multiple granular balls and higher classification accuracy on all three datasets.',
    'We varied one parameter at a time':
        'We varied one parameter at a time on Cora, CiteSeer, and CoraFull under uniform 0.4 noise, keeping the pretrained representations fixed and using five paired seeds (Appendix I.1). CoraFull was sensitive to the ball budget: reducing $B_max$ from 64 to 32 lowered mean accuracy by 2.10 percentage points, whereas increasing it to 128 improved accuracy by 1.57 points. The higher budget improved every seed on CoraFull. By contrast, Cora still formed 52\u201353 balls with a budget of 128, leaving its targets unchanged. Increasing the majority-class weight $\u03b2$ from 0.5 to 1 improved CiteSeer by 1.74 points, with little change on the other two datasets.',
    'GB-LCSS uses labels from related':
        'GB-LCSS supplements an individual annotation with labels from related training nodes. Label-free grouping keeps corrupted annotations from shaping the groups, while excluding the current node prevents its label from reinforcing the consensus used to adjust it. Proposition 1 quantifies the reliability of the local vote when the true class is more frequent than any competing class among the voting nodes and labels are corrupted independently by symmetric noise. The random-grouping and smoothing controls (Table 5) support the value of selecting informative local evidence. The target-error decomposition in Appendix I.2 further shows how the targets improved: the reduction in error on corrupted labels outweighed the increase on initially correct labels in the tested runs.',
    'Granular balls organize local-consensus':
        'Granular balls organize this supervision into shared regions. Each node is assigned to one group, class statistics are reused within that group, and the prescribed budget bounds the number of regions. The sensitivity results in Section 4.3.4 show why local grouping and voting coverage matter: the largest losses occurred when all nodes remained in one ball or substantially fewer nodes received consensus supervision. In Table 6, granular ball grouping achieves accuracy within one percentage point of both kNN and k-means on four of the five datasets; on CoraFull, kNN is higher by 3.11 points. Granular ball grouping also uses less storage than exact kNN on the three larger datasets and has the shortest construction time on three of the five datasets in the reported measurements (Fig. 5). Together, these results support granular balls as a practical choice for combining competitive accuracy with shared regional statistics and favorable grouping costs.',
}

SUPP_REWRITES = {
    'Targets were computed on the CPU':
        'Targets were computed on the CPU for these reruns, whereas the main comparison used GPU target construction. The MLP was trained on an AMD Radeon RX 7700 XT using PyTorch 2.9.1+rocm7.2.1. The scikit-learn 1.8.0 implementation of 2-means used six CPU threads; PyTorch used four host threads. The default configuration was rerun alongside the controls and parameter variants. Its results in Table 7 and Table I.1 serve as the reference for these comparisons and differ slightly from those in the main comparison.',
    'We also examined the local-class condition':
        "We also examined the local-class condition in Proposition 1. Using clean labels for this diagnostic, we counted training nodes whose true class had more votes than any other class among the other training nodes in their ball. Averaged over seeds, these fractions were 75.64%, 65.02%, and 36.13% on Cora, CiteSeer, and CoraFull, respectively. They describe where the proposition's positive-margin assumption holds. Table I.2 measures a different quantity: the net change in soft-target error across all training nodes.",
}


def update_references(doc):
    for node in doc.element.xpath('.//w:p'):
        from docx.text.paragraph import Paragraph
        p = Paragraph(node, doc._body)
        joined = ''.join(node.xpath('.//w:t/text()'))
        replacements = []
        for match in re.finditer(r'\((1[4-9]|2[0-3])\)', joined):
            old = match.group(0)
            replacements.append((old, '(' + str(EQ_MAP[int(match.group(1))]) + ')'))
        for match in re.finditer(r'Section (3\.[456])\b', joined):
            replacements.append((match.group(0), 'Section ' + SECTION_MAP[match.group(1)]))
        # Temporary tokens prevent a new equation number being mapped a second time.
        for i, (old, new) in enumerate(dict(replacements).items()):
            replace_in_paragraph(p, old, f'REFPLACEHOLDER{i}')
        for i, (old, new) in enumerate(dict(replacements).items()):
            replace_in_paragraph(p, f'REFPLACEHOLDER{i}', new)


def math_signature(node):
    node = deepcopy(node)
    # Reverse only the agreed notational change for equation-preservation checks.
    for sub in list(node.iter(qn('m:sSubSup'))):
        if ''.join(t.text or '' for t in sub.find(qn('m:sup')).iter(qn('m:t'))) != 'clean':
            continue
        old = OxmlElement('m:sSub')
        base = deepcopy(sub.find(qn('m:e')))
        next(base.iter(qn('m:t'))).text = '\u03c0'
        old.extend([base, deepcopy(sub.find(qn('m:sub')))])
        sub.getparent().replace(sub, old)
    for e in list(node.iter()):
        if e.tag in {qn('w:rPr'), qn('m:rPr')}:
            e.getparent().remove(e)
    return etree.tostring(node, method='c14n', exclusive=True)


def verify(source, output, is_main):
    old, new = Document(source), Document(output)
    assert len(old.tables) == len(new.tables)
    for i, (a, b) in enumerate(zip(old.tables, new.tables)):
        assert (len(a.rows), len(a.columns)) == (len(b.rows), len(b.columns))
        for row_a, row_b in zip(a.rows, b.rows):
            for cell_a, cell_b in zip(row_a.cells, row_b.cells):
                ta = '\n'.join(text(p) for p in cell_a.paragraphs)
                tb = '\n'.join(text(p) for p in cell_b.paragraphs)
                if not is_main and i == 9:
                    ta = ta.replace('Eq. (14)', 'Eq. (17)')
                if not is_main and i == 15:
                    ta = ta.replace('Cost oncorrect labels', 'Error increase oncorrect labels')
                assert ta == tb, (i, ta, tb)
    with ZipFile(source) as a, ZipFile(output) as b:
        media = [n for n in a.namelist() if n.startswith('word/media/')]
        assert all(a.read(n) == b.read(n) for n in media)
    old_refs = [text(p) for p in old.paragraphs if re.match(r'^\[\d+\]', text(p))]
    new_refs = [text(p) for p in new.paragraphs if re.match(r'^\[\d+\]', text(p))]
    assert old_refs == new_refs
    # Displayed equations are unchanged except for notation and display order.
    def equation_math(doc):
        return Counter(math_signature(m) for p in doc.paragraphs
                       if re.search(r'\((?:\d+|H\.\d+)\)$', text(p))
                       for m in p._p.xpath('./m:oMath | ./m:oMathPara'))
    assert equation_math(old) == equation_math(new), 'Displayed mathematics changed'
    if is_main:
        numbers = [int(re.search(r'\((\d+)\)$', text(p)).group(1))
                   for p in new.paragraphs if p._p.xpath('.//m:oMath') and re.search(r'\((\d+)\)$', text(p))]
        assert numbers == list(range(1, 24)), numbers
        titles = [text(p) for p in new.paragraphs if re.match(r'^3\.[456] ', text(p))]
        assert titles == ['3.4 Analysis of local-consensus supervision',
                          '3.5 Graph-teacher distillation extension', '3.6 Computational costs'], titles
        for i in [100]:
            assert etree.tostring(old.paragraphs[i]._p) in [etree.tostring(p._p) for p in new.paragraphs]
    assert 'REFPLACEHOLDER' not in '\n'.join(text(p) for p in new.paragraphs)
    return {'output': str(output), 'tables': len(new.tables), 'table_results_unchanged': True,
            'displayed_math_preserved_except_clean_proportion_notation': True,
            'equation_numbers_and_section_references_updated': True,
            'images_unchanged': True, 'references_unchanged': True}


def build(is_main):
    stem = 'GB-LCSS' if is_main else 'GB-LCSS_supplementary'
    source = DOWNLOADS / (stem + '_experiments_polished_20260920_red.docx')
    output = DOWNLOADS / (stem + '_theory_experiments_integrated_20260920_red.docx')
    old_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    doc = Document(source)
    before = [(p._p, text(p)) for p in doc.paragraphs]
    rewrites = MAIN_REWRITES if is_main else SUPP_REWRITES
    for prefix, replacement in rewrites.items():
        rewrite(find(doc, prefix), replacement)
    if is_main:
        p = find(doc, 'For a training node v, let')
        replace_in_paragraph(p, 'For a training node ', 'Consider a training node ')
        replace_in_paragraph(p, ', let ', '. For analysis, let ')
        bound = find(doc, 'The event includes a tie')
        replace_in_paragraph(bound,
            'For a fixed positive margin, more independent votes strengthen the concentration around the correct class.',
            'For fixed class count, noise rate, and positive margin, the exponential term in the bound decreases as the number of independent votes increases.')
        theory_start = next(i for i,p in enumerate(doc.paragraphs) if text(p).startswith('3.6 Analysis'))
        theory_end = next(i for i,p in enumerate(doc.paragraphs) if text(p) == '4 Experiments')
        theory = doc.paragraphs[theory_start:theory_end]
        assert rename_clean_proportions(theory) == 4
        teacher = find(doc, '3.4 Graph-teacher')
        for p in theory:
            teacher._p.addprevious(p._p)
        rewrite(find(doc, '3.6 Analysis'), '3.4 Analysis of local-consensus supervision')
        rewrite(find(doc, '3.4 Graph-teacher'), '3.5 Graph-teacher distillation extension')
        rewrite(find(doc, '3.5 Computational'), '3.6 Computational costs')
    else:
        theory = [p for i,p in enumerate(doc.paragraphs) if 45 <= i < 69]
        assert rename_clean_proportions(theory) == 8
        table = doc.tables[14]
        for index in [2, 6, 8, 10, 13, 15]:
            for cell in table.rows[index].cells:
                pr = cell._tc.get_or_add_tcPr()
                borders = pr.find(qn('w:tcBorders'))
                if borders is None:
                    borders = OxmlElement('w:tcBorders')
                    pr.append(borders)
                top = OxmlElement('w:top')
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'BFBFBF')]:
                    top.set(qn('w:' + key), value)
                borders.append(top)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
        rewrite(doc.tables[15].cell(0,4).paragraphs[0], 'Error increase on\ncorrect labels')
    update_references(doc)
    doc.save(output)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == old_hash
    CHANGES.extend({'old': old, 'new': text(p)} for node, old in before
                   for p in doc.paragraphs if p._p is node and old != text(p))
    return verify(source, output, is_main)


if __name__ == '__main__':
    checks = [build(True), build(False)]
    (HERE / 'checks.json').write_text(json.dumps(checks, indent=2), encoding='utf8')
    (HERE / 'changes.json').write_text(json.dumps(CHANGES, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps(checks, indent=2))
