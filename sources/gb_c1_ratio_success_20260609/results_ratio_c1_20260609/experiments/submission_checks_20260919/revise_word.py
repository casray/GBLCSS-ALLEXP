"""Integrate verified supplementary experiments without changing the main results."""
from collections import Counter
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree


HERE = Path(__file__).resolve().parent
DOWNLOADS = Path('C:/Users/Admin/Downloads')
MAIN_IN = DOWNLOADS/'GB-LCSS_theory_revised_20260918.docx'
SUPP_IN = DOWNLOADS/'GB-LCSS_supplementary_theory_20260918.docx'
MAIN_OUT = DOWNLOADS/'GB-LCSS_experiments_revised_20260919.docx'
SUPP_OUT = DOWNLOADS/'GB-LCSS_supplementary_experiments_20260919.docx'
RESULTS = HERE/'gpu_reference_matched'
summary = {(r['dataset'], r['arm']): r for r in csv.DictReader((RESULTS/'summary_verified.csv').open(encoding='utf-8-sig'))}
theory = {r['dataset']:r for r in csv.DictReader((RESULTS/'theory_diagnostics.csv').open(encoding='utf-8-sig'))}
DATASETS = ['cora','citeseer','corafull']
NAMES = {'cora':'Cora','citeseer':'CiteSeer','corafull':'CoraFull'}


def content(p):
    return ''.join(p._p.xpath('.//w:t | .//m:t')[i].text or ''
                   for i in range(len(p._p.xpath('.//w:t | .//m:t'))))


def find(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix,len(matches))
    return matches[0]


def plain(p, text):
    assert not p._p.xpath('.//m:oMath | .//w:drawing'), p.text
    for child in list(p._p):
        if child.tag not in {qn('w:pPr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')}:
            p._p.remove(child)
    r=p.add_run(text)
    r.font.color.rgb=RGBColor(0,0,0)
    return p


def replace_text(p, old, new):
    matches = [t for t in p._p.xpath('.//w:t') if old in (t.text or '')]
    assert len(matches)==1, (old,p.text)
    matches[0].text=matches[0].text.replace(old,new)


def add_after(anchor, text, template=None, style=None):
    node=OxmlElement('w:p')
    if template is not None and template._p.pPr is not None:
        node.append(deepcopy(template._p.pPr))
    anchor._p.addnext(node)
    p=Paragraph(node,anchor._parent)
    if style is not None:
        p.style=style
    p.add_run(text).font.color.rgb=RGBColor(0,0,0)
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.page_break_before=False
    return p


def math_run(value, upright=False):
    r=OxmlElement('m:r')
    pr=OxmlElement('m:rPr')
    sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p' if upright else 'i');pr.append(sty);r.append(pr)
    t=OxmlElement('m:t');t.text=value;r.append(t)
    return r


def inline_math(text):
    om=OxmlElement('m:oMath')
    for token in re.split(r'(B_max|n_min|s_min|A\^2X)',text):
        if not token: continue
        if token in {'B_max','n_min','s_min'}:
            base, sub=token.split('_')
            s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base));u=OxmlElement('m:sub');u.append(math_run(sub,True));s.extend([e,u]);om.append(s)
        elif token=='A^2X':
            s=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(math_run('A'));u=OxmlElement('m:sup');u.append(math_run('2',True));s.extend([e,u]);om.append(s);om.append(math_run('X'))
        else:
            om.append(math_run(token))
    return om


def typeset(p):
    # Dollar-delimited short symbols become editable Word mathematics.
    text=p.text
    if '$' not in text:return
    plain(p,'')
    for i,part in enumerate(text.split('$')):
        if not part:continue
        if i%2:p._p.append(inline_math(part))
        else:p.add_run(part).font.color.rgb=RGBColor(0,0,0)


def accuracy(dataset,arm):
    r=summary[dataset,arm]
    return f"{float(r['accuracy_percent']):.2f} ± {float(r['sd_percent']):.2f}"


def new_table(doc,headers,rows,widths,font_size=10):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single' if name in {'top','bottom'} else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
    pr.append(borders)
    t._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),str(round(sum(widths)*1440)))
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for row in rows:t.add_row()
    for i,values in enumerate([headers,*rows]):
        tr=t.rows[i]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value,width) in enumerate(zip(t.rows[i].cells,values,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for name,w in [('top',45),('bottom',45),('left',55),('right',55)]:
                e=OxmlElement('w:'+name);e.set(qn('w:w'),str(w));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            if i==0:
                b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');b.append(e);cp.append(b)
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=(i==0)
            r=p.add_run(value);r.font.name='Times New Roman';r.font.size=Pt(font_size);r.font.bold=(i==0);r.font.color.rgb=RGBColor(0,0,0)
            typeset(p)
            for rr in p.runs:rr.font.size=Pt(font_size);rr.font.name='Times New Roman';rr.font.bold=(i==0)
    return t


def append_paragraph(doc,text,style='Body Text'):
    p=doc.add_paragraph(text,style=style)
    p.paragraph_format.keep_with_next=style.startswith('Heading') or style=='Caption'
    p.paragraph_format.widow_control=True
    p.paragraph_format.page_break_before=False
    if style == 'Caption':
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    typeset(p)
    return p


def edit_main():
    d=Document(MAIN_IN)
    replace_text(find(d,'We evaluate GB-LCSS on five'),'Appendices A–H','Appendices A–I')
    p=next(p for p in d.paragraphs if p.text.startswith('GB-LCSS uses ') and 'throughout the experiments.' in p.text)
    replace_text(p,'GB-LCSS uses ','Unless varied in Section 4.3.4, GB-LCSS uses ')
    replace_text(p,'throughout the experiments.','as the default settings.')
    replace_text(find(d,'Equation (23) holds'),'The Brier analysis in Section 4.3.1 measures the aggregate target error empirically.',
                 'Section 4.3.1 and Appendix I.2 examine this balance empirically.')

    anchor=find(d,'Fig. 4. Brier score')
    add_after(anchor,
        'The supplementary runs under uniform 0.4 separate this reduction into two components: lower error on corrupted labels and increased error on initially correct labels (Appendix I.2). The former exceeds the latter in all 15 dataset–seed runs. Random grouping also lowers Brier scores, but yields much lower classification accuracy, showing that target error and downstream performance capture different aspects of supervision.',
        template=find(d,'Under the same construction settings'))

    p=find(d,'The representation used for grouping')
    plain(p,'The representation used for grouping determines which training nodes contribute labels to each target. Table 7 compares the frozen GRACE representation with raw attributes $X$ and two-step propagated attributes $A^2X$ on Cora, CiteSeer, and CoraFull under uniform 0.4 noise. The propagated representation uses the original adjacency matrix without added self-loops or adjacency normalization; the resulting node vectors are L2-normalized as in Eq. (3). All three conditions were rerun over five paired seeds using the setup in Appendix I.1; ball counts are reported as ranges across those seeds.')
    p.paragraph_format.keep_with_next=False;typeset(p)
    plain(find(d,'Table 7. Classification'),'Table 7. Classification accuracy and ball counts for different representations under uniform 0.4 noise (five paired seeds).')
    p=find(d,'Under the same construction settings')
    plain(p,'Under the same construction settings, neither raw attributes nor propagated attributes yield an accepted split in any of the 15 dataset–seed runs. All nodes therefore remain in one granular ball. For each training node, the two representations give the same consensus from all other training nodes, the same soft target, and identical accuracy under the paired training procedure. GRACE produces multiple balls, allowing each target to draw on a smaller local group. The accuracy differences in Table 7 show that the representation affects whether the fixed splitting rule forms useful regions for supervision.')
    p.paragraph_format.space_before=Pt(5)

    old=d.tables[6]
    rows=[]
    for dataset in DATASETS:
        r=summary[dataset,'default'];lo=int(float(r['balls_min']));hi=int(float(r['balls_max']))
        rows.append([NAMES[dataset],accuracy(dataset,'default'),accuracy(dataset,'raw_x'),accuracy(dataset,'a2x'),
                     str(lo) if lo==hi else f'{lo}–{hi}','1'])
    new=new_table(d,['Dataset','GRACE','Raw X','A²X','Balls from\nGRACE','Balls from\nX / A²X'],rows,[.88,1.23,1.23,1.23,1.04,1.04],9.5)
    old._tbl.addprevious(new._tbl);old._tbl.getparent().remove(old._tbl)

    anchor=find(d,'We also examine pretraining variability')
    h=add_after(anchor,'4.3.4 Sensitivity to construction and target parameters',style='Heading 4')
    h.paragraph_format.keep_with_next=True
    p=add_after(h,'We varied one parameter at a time on Cora, CiteSeer, and CoraFull under uniform 0.4 noise, using five paired seeds and fixed pretrained representations (Appendix I.1). Increasing $B_max$ from 64 to 128 changed mean accuracy by 0.00, 0.21, and 1.57 percentage points, respectively. Cora still formed 52–53 balls, whereas CoraFull reached 128 balls and improved on all five seeds. Increasing $β$ from 0.5 to 1 improved CiteSeer by 1.74 points but gave little average benefit on the other two datasets.',template=p)
    typeset(p)
    p=add_after(p,'The support and split thresholds have clearer failure cases. Raising $s_min$ from 2 to 10 reduced the fraction of training nodes receiving consensus correction to 56.53% on Cora and 58.32% on CiteSeer, with accuracy decreases of 7.27 and 5.22 points. Setting $δ=0.10$ prevented all splits on the three datasets, whereas $δ=0.01$ produced the same targets as the default 0.05. The effect of $n_min$ differed across datasets. These results show that region resolution and voting coverage both matter; the default values provide a common configuration rather than a dataset-wise optimum.',template=p)
    typeset(p)
    d.save(MAIN_OUT)


def edit_supp():
    d=Document(SUPP_IN)
    h=append_paragraph(d,'Appendix I. Parameter sensitivity and target-error decomposition','Heading 2')
    h.paragraph_format.page_break_before=True
    append_paragraph(d,'I.1 Parameter sensitivity and representation controls','Heading 3')
    append_paragraph(d,'We evaluated Cora, CiteSeer, and CoraFull under uniform 0.4 training-label noise, with the 20%/10%/70% split, clean validation labels, and seeds 0–4. Each dataset used the same frozen GRACE representation as the main comparisons. The target rule, MLP architecture, optimizer, 200-epoch budget, and validation-based checkpoint selection were unchanged. Each alternative was compared with a default run using the same split, noise realization, and initialization seed.')
    append_paragraph(d,'For these runs, targets were computed on the CPU and MLP training used an AMD Radeon RX 7700 XT with PyTorch 2.9.1+rocm7.2.1. The sklearn 1.8.0 two-means routine used six CPU threads; PyTorch used four host threads. Table 7 and Table I.1 report this paired rerun. All parameter differences below are measured against its default, not against a result from a different numerical execution path.')
    append_paragraph(d,'The default configuration is $β=0.5$, $B_max=64$, $n_min=32$, $δ=0.05$, and $s_min=2$. Each sensitivity row changes only the named parameter. Here, $s_min$ counts votes remaining after exclusion of the current node. Setting $β=0$ retains the empirical local distribution without extra majority-class sharpening. Random grouping preserves the default group sizes; raw $X$ and $A^2X$ change the representation used for grouping, as described in Section 4.3.3.')
    append_paragraph(d,'Table I.1. Parameter sensitivity and controls under uniform 0.4 noise. Accuracy is mean ± sample standard deviation over five paired seeds.','Caption')
    labels=[('default','Default'),('beta_0','$β=0$'),('beta_0.25','$β=0.25$'),('beta_0.75','$β=0.75$'),('beta_1','$β=1$'),
            ('max_balls_32','$B_max=32$'),('max_balls_128','$B_max=128$'),
            ('min_train_per_ball_5','$s_min=5$'),('min_train_per_ball_10','$s_min=10$'),
            ('min_ball_size_1','$n_min=1$'),('min_ball_size_16','$n_min=16$'),('min_ball_size_64','$n_min=64$'),
            ('split_improvement_0.01','$δ=0.01$'),('split_improvement_0.1','$δ=0.10$'),
            ('mlp_ce','MLP+CE'),('shuffled','Random size-matched grouping'),('raw_x','Raw $X$'),('a2x','$A^2X$')]
    rows=[[label,*[accuracy(ds,arm) for ds in DATASETS]] for arm,label in labels]
    new_table(d,['Configuration','Cora','CiteSeer','CoraFull'],rows,[2.20,1.4833,1.4833,1.4834],10)
    append_paragraph(d,'When two configurations produced exactly the same target tensor for a given dataset and seed, they shared the same training result. Thus, the 270 configuration–dataset–seed combinations required 214 training runs and 56 exact-target reuses. Equality of the targets was checked directly; the reused entries are not additional independent observations.')
    append_paragraph(d,'The effect of the ball budget depends on whether it limits the partition. Cora formed 52–53 balls with both budgets of 64 and 128. CiteSeer formed 63–64 and 63–66 balls, respectively. CoraFull reached the budget in both cases: increasing it to 128 raised accuracy on all five seeds, with a paired gain of 1.57 ± 0.54 percentage points, and reduced mean target Brier score from 0.6321 to 0.6026. This is consistent with the benefit of finer regions on that dataset.')
    append_paragraph(d,'Increasing $β$ to 1 improved CiteSeer on all five seeds, by 1.74 ± 0.53 points in paired differences, but the mean changes on Cora and CoraFull were only +0.04 and −0.14 points. Reducing $n_min$ to 1 changed their means by −1.65, +0.70, and −0.05 points, respectively. Neither adjustment gives a uniform gain across the three datasets.')
    append_paragraph(d,'A high support requirement reduces the coverage of consensus correction. With $s_min=10$, coverage was 56.53%, 58.32%, and 99.13% on Cora, CiteSeer, and CoraFull, respectively; the corresponding accuracy changes were −7.27, −5.22, and −0.10 points. A stricter split threshold, $δ=0.10$, left all nodes in one ball on every seed. Its targets were identical to those obtained from raw $X$ and $A^2X$. At $δ=0.01$, the targets instead matched the default $δ=0.05$ in all 15 runs. The full results distinguish stable settings from settings that prevent local regions from forming.')
    append_paragraph(d,'I.2 Error reduction on training targets','Heading 3')
    append_paragraph(d,'Table I.2 evaluates the decomposition in Eq. (H.5) on the default runs. For each seed, the reduction on corrupted labels is the sum of their one-hot errors minus their soft-target errors, divided by the total number of training nodes. The cost on correct labels is the sum of their soft-target errors, divided by the same denominator; their original one-hot error is zero. The net reduction is the former minus the latter. We then average each quantity over five seeds. Clean training labels are used only for these measurements.')
    append_paragraph(d,'Table I.2. Decomposition of the reduction in training-target Brier score under uniform 0.4 noise.','Caption')
    keys=['noisy_brier','target_brier','corrupted_gain','clean_cost','net_gain']
    rows=[[NAMES[ds],*[f'{float(theory[ds][k]):.4f}' for k in keys]] for ds in DATASETS]
    new_table(d,['Dataset','One-hot\nlabel','Soft\ntarget','Reduction on\ncorrupted labels','Cost on\ncorrect labels','Net\nreduction'],rows,[.90,.90,.90,1.55,1.50,.90],9.5)
    p=append_paragraph(d,'The reduction on corrupted labels exceeds the cost on correct labels in every dataset–seed run. The positive net values therefore reflect a balance of these two effects, rather than preservation of every correct annotation. Random grouping also reduces Brier scores relative to observed one-hot labels, yet its classification accuracy remains much lower than that of the default regions (Table I.1). Lower target error alone is therefore insufficient to explain the benefit of informative group membership.')
    p.paragraph_format.space_before=Pt(5)
    append_paragraph(d,'The positive-margin condition in Proposition 1 holds for 75.64%, 65.02%, and 36.13% of training nodes on Cora, CiteSeer, and CoraFull, respectively, when measured from the clean labels of the other training nodes in each ball and averaged across seeds. These are empirical node fractions, not estimates of the probability bound. They show why the condition must be stated explicitly: a ball can provide useful evidence for some members without its clean majority matching every member’s class.')
    d.save(SUPP_OUT)


def math_counter(d):
    return Counter(hashlib.sha256(etree.tostring(x,method='c14n',exclusive=True)).hexdigest()
                   for x in d.element.xpath('.//m:oMath'))


def verify():
    checks={}
    for label,source,target in [('main',MAIN_IN,MAIN_OUT),('supplement',SUPP_IN,SUPP_OUT)]:
        old,new=Document(source),Document(target)
        assert not (math_counter(old)-math_counter(new)), 'original mathematical object changed'
        with ZipFile(source) as a, ZipFile(target) as b:
            media=[n for n in a.namelist() if n.startswith('word/media/')]
            assert all(a.read(n)==b.read(n) for n in media)
        if label=='main':
            for i,(a,b) in enumerate(zip(old.tables,new.tables)):
                if i!=6:
                    assert etree.tostring(a._tbl,method='c14n',exclusive=True)==etree.tostring(b._tbl,method='c14n',exclusive=True),i
            oldrefs=[content(p) for p in old.paragraphs[222:]]
            idx=next(i for i,p in enumerate(new.paragraphs) if p.text=='References')
            assert oldrefs==[content(p) for p in new.paragraphs[idx+1:]]
            assert new.paragraphs[0].text==old.paragraphs[0].text
        else:
            for a,b in zip(old.tables,new.tables):
                assert etree.tostring(a._tbl,method='c14n',exclusive=True)==etree.tostring(b._tbl,method='c14n',exclusive=True)
        checks[label]={'input':str(source),'output':str(target),'tables_before':len(old.tables),
                       'tables_after':len(new.tables),'original_math_objects_preserved':True,
                       'media_preserved':len(media),'reference_content_preserved':True}
    (HERE/'word_revision_checks.json').write_text(json.dumps(checks,indent=2),encoding='utf8')
    print(json.dumps(checks,indent=2))


if __name__=='__main__':
    edit_main();edit_supp();verify()
