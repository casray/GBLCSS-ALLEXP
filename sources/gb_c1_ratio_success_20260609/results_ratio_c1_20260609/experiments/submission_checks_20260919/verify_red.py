from zipfile import ZipFile

from docx import Document


ROOT = "C:/Users/Admin/Downloads/"
PAIRS = [
    ("GB-LCSS_experiments_revised_20260919.docx", "GB-LCSS_experiments_revised_20260919_red.docx"),
    ("GB-LCSS_supplementary_experiments_20260919.docx", "GB-LCSS_supplementary_experiments_20260919_red.docx"),
]


def colors(doc):
    return {
        "red": len(doc.element.xpath('.//w:color[@w:val="FF0000"]')),
        "black": len(doc.element.xpath('.//w:color[@w:val="000000"]')),
        "math_red": len(doc.element.xpath('.//m:color[@m:val="FF0000"]')),
    }


for source_name, marked_name in PAIRS:
    source = Document(ROOT + source_name)
    marked = Document(ROOT + marked_name)
    assert [p.text for p in source.paragraphs] == [p.text for p in marked.paragraphs]
    assert len(source.tables) == len(marked.tables)
    for source_table, marked_table in zip(source.tables, marked.tables):
        source_cells = [[cell.text for cell in row.cells] for row in source_table.rows]
        marked_cells = [[cell.text for cell in row.cells] for row in marked_table.rows]
        assert source_cells == marked_cells
    with ZipFile(ROOT + source_name) as source_zip, ZipFile(ROOT + marked_name) as marked_zip:
        source_media = sorted(n for n in source_zip.namelist() if n.startswith("word/media/"))
        marked_media = sorted(n for n in marked_zip.namelist() if n.startswith("word/media/"))
        assert source_media == marked_media
        assert all(source_zip.read(n) == marked_zip.read(n) for n in source_media)
    print(source_name, "content and media preserved", colors(marked))
